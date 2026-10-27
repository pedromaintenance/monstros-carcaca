#!/usr/bin/env python3
"""Extração fiel dos movimentos do Bárbaro, com REQUER como relação obrigatória.
Apenas anota relações evidentes; não autoriza seleção automática nem altera cânone.
Uso: python extrair_barbaro_docx.py <docx> <inventario.json> <saida.json>
"""
from pathlib import Path
from docx import Document
import hashlib, json, sys, unicodedata, re

def slug(t):
    t=unicodedata.normalize('NFKD',t).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',t).strip('-')

def extract(docx_path,inventory_path):
    docx_path=Path(docx_path)
    inventory=json.loads(Path(inventory_path).read_text(encoding='utf8'))
    h=hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if h!=inventory['fonte_sha256']:raise ValueError('Hash da fonte diferente do inventário')
    doc=Document(docx_path)
    rows=[m for m in inventory['movimentos'] if m['grupo']=='BÁRBARO']
    if len(rows)!=59 or rows[0]['paragrafo_docx']!=1293 or rows[-1]['paragrafo_docx']!=1447:
        raise ValueError('Âncoras ou quantidade do Bárbaro alteradas')
    all_ids={m['id']:m for m in inventory['movimentos']}
    edges={m['origem_movimento_id']:m for m in inventory['dependencias_exatas_candidatas']}
    unresolved={m['origem_movimento_id']:m for m in inventory['dependencias_nao_resolvidas']}
    title_map={slug(m['titulo']):m for m in rows}
    movements=[]
    for k,item in enumerate(rows):
        i=item['paragrafo_docx'];stop=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else 1453
        p=doc.paragraphs[i]
        lines=p.text.strip().splitlines()
        name=lines.pop(0).strip()
        if name!=item['titulo']:raise ValueError(f'Título conflitante #{i}: {name}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).strip().split(':',1)[1].strip()
        if req!=item['requisito_literal']:
            raise ValueError(f'REQUER conflitante #{i}: {req} != {item["requisito_literal"]}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError('Corpo ausente no movimento '+name)
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}]
        opts=[]
        for j in range(i+1,stop):
            other=doc.paragraphs[j]
            if other.style.name.startswith('Heading'):
                raise ValueError('Cabeçalho inesperado no movimento '+name+' em '+str(j))
            text=other.text.strip()
            if not text:continue
            typ='opcao' if other.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':text})
            if typ=='opcao':opts.append({'paragrafo_docx':j,'texto':text})
        exact=edges.get(item['id']);pending=unresolved.get(item['id'])
        if req:
            if bool(exact)==bool(pending):raise ValueError('Vínculo múltiplo ou ausente '+name)
            if exact:
                target=all_ids[exact['requisito_movimento_id']]
                if slug(req)!=slug(target['titulo']):raise ValueError('REQ não corresponde literalmente '+name)
                structured={'tipo':'movimento_exato','movimento_id':target['id'],
                    'titulo_requerido':target['titulo'],'escopo':exact['escopo'],
                    'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            else:
                if i!=1430 or req!='Fúria Divina e Demasiado Furioso Para Morrer':
                    raise ValueError(f'REQUER especial não previsto #{i}: {req}')
                required_names=['Fúria Divina','Demasiado Furioso Para Morrer']
                structured={'tipo':'todos','itens':[
                    {'tipo':'movimento_exato','movimento_id':title_map[slug(name)]['id'],
                     'titulo_requerido':name} for name in required_names],
                    'operador_literal':'e',
                    'status':'interpretacao_literal_pendente_homologacao'}
        else:structured=None
        movements.append({'id':item['id'],'paragrafo_docx':i,'grupo':'BÁRBARO',
            'categoria':'classe','titulo':name,'requisito_literal':req,
            'requisito_estruturado':structured,'partes':parts,'opcoes':opts,
            'texto_original':'\n'.join(x['texto'] for x in parts),
            'texto_integral_conferido':True,'estado':'transcricao_integral_sem_revisao_mecanica',
            'aplicacao_canonica':False,'origem_inventario_id':item['id']})
    if movements[-1]['titulo']!='LAMENTO':raise ValueError('Último movimento incorreto')
    if len(movements)!=59:raise ValueError('Total incorreto')
    if sum(bool(m['requisito_literal']) for m in movements)!=37:raise ValueError('Total REQUER incorreto')
    if sum(m['requisito_estruturado']['tipo']=='movimento_exato' for m in movements if m['requisito_estruturado'])!=36:
        raise ValueError('Dependências exatas incorretas')
    result={'schema_version':'0.1.0','fonte':inventory['fonte'],'fonte_sha256':h,
      'grupo':'BÁRBARO','posicao':'indices_paragrafo_docx_zero_based',
      'descricao':'Transcrição completa do Bárbaro. REQUER é obrigatório; correspondências automáticas não são homologação.',
      'movimentos':movements,
      'integridade':{'total_movimentos':len(movements),
          'total_com_requer':sum(bool(m['requisito_literal']) for m in movements),
          'total_vinculos_exatos':36,'total_requer_composto_ou_condicional':1,
          'total_opcoes':sum(len(m['opcoes']) for m in movements),
          'total_partes':sum(len(m['partes']) for m in movements)}}
    return result

if __name__=='__main__':
    d=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inv=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    dest=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_barbaro/hack-carcaca-dw-barbaro-integrais.json')
    result=extract(d,inv)
    dest.parent.mkdir(parents=True,exist_ok=True)
    dest.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(result['integridade'],ensure_ascii=False,indent=2))