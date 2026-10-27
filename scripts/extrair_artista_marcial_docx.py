#!/usr/bin/env python3
"""Extrai movimentos integrais de ARTISTA MARCIAL do Hack Carcaça.

python extrair_artista_marcial_docx.py MANUSCRITO.docx INVENTARIO.json SAIDA.json
Preserva REQUER exato e a conjunção obrigatória de VOO sem promoção canônica.
"""
from __future__ import annotations
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path
from docx import Document

FIRST, LAST, NEXT_HEADING = 1862, 1963, 1966
SPECIAL_AND = {1962: ('Forma do Ar', 'Corpo Sem Peso')}


def slug(text: str) -> str:
    return re.sub(r'[^a-z0-9]+', '-', unicodedata.normalize('NFKD', text)
                  .encode('ascii', 'ignore').decode().lower()).strip('-')


def extract(docx_path, inventory_path):
    docx_path, inventory_path = Path(docx_path), Path(inventory_path)
    inventory = json.loads(inventory_path.read_text(encoding='utf-8'))
    source_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if source_hash != inventory['fonte_sha256']:
        raise ValueError('Manuscrito não corresponde ao SHA256 do inventário')
    doc = Document(docx_path)
    rows = [m for m in inventory['movimentos'] if m['grupo']=='ARTISTA MARCIAL']
    if len(rows)!=42 or rows[0]['paragrafo_docx']!=FIRST or rows[-1]['paragrafo_docx']!=LAST:
        raise ValueError('Inventário Artista Marcial alterado: 42 movimentos esperados')
    if (doc.paragraphs[1860].text.strip()!='ARTISTA MARCIAL'
            or doc.paragraphs[NEXT_HEADING].text.strip()!='PALADINO'
            or not doc.paragraphs[NEXT_HEADING].style.name.startswith('Heading')):
        raise ValueError('Limites entre as classes não conferem com a fonte')
    by_title = {}
    for m in rows: by_title.setdefault(slug(m['titulo']), []).append(m)
    all_by_id={m['id']:m for m in inventory['movimentos']}
    edges={m['origem_movimento_id']:m for m in inventory['dependencias_exatas_candidatas']}
    pending={m['origem_movimento_id']:m for m in inventory['dependencias_nao_resolvidas']}
    out=[]; occupied=set()
    for k, original in enumerate(rows):
        i=original['paragrafo_docx']
        stop=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else NEXT_HEADING
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=original['titulo']:
            raise ValueError(f'Cabeçalho alterado no parágrafo {i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=original['requisito_literal']:
            raise ValueError(f'REQUER mudou no DOCX #{i}')
        body='\n'.join(lines).strip()
        if not body:
            raise ValueError(f'Corpo vazio no movimento #{i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}]; opts=[]
        if i in occupied: raise ValueError(f'Parágrafo duplicado: {i}')
        occupied.add(i)
        for j in range(i+1,stop):
            p=doc.paragraphs[j]
            if p.style.name.startswith('Heading'):
                raise ValueError(f'Cabeçalho inesperado dentro de movimento: {j}')
            t=p.text.strip()
            if not t:continue
            if j in occupied:raise ValueError(f'Parágrafo repetido: {j}')
            occupied.add(j)
            typ='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':t})
            if typ=='opcao':opts.append({'paragrafo_docx':j,'texto':t})
        direct=edges.get(original['id']); unresolved=pending.get(original['id'])
        if req:
            if bool(direct)==bool(unresolved):
                raise ValueError(f'Requisito não exclusivo no inventário: {i}')
            if direct:
                target=all_by_id[direct['requisito_movimento_id']]
                if slug(req)!=slug(target['titulo']):
                    raise ValueError(f'Referência sem título literal correspondente: {i}')
                requirement={'tipo':'movimento_exato','movimento_id':target['id'],
                            'titulo_requerido':target['titulo'],'escopo':direct['escopo'],
                            'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            elif i in SPECIAL_AND:
                target_names=SPECIAL_AND[i]
                if req!=' e '.join(target_names) or unresolved['requisito_literal']!=req:
                    raise ValueError(f'Conjunção de VOO divergente: {req}')
                targets=[]
                for name in target_names:
                    candidates=by_title.get(slug(name),[])
                    if len(candidates)!=1:raise ValueError(f'Alvo ambíguo/ausente: {name}')
                    t=candidates[0]
                    targets.append({'tipo':'movimento_exato','movimento_id':t['id'],
                                    'titulo_requerido':t['titulo']})
                requirement={'tipo':'todos','itens':targets,'operador_literal':'e',
                    'status':'interpretacao_literal_pendente_homologacao'}
            else:
                raise ValueError(f'REQUER especial não mapeado: {i}: {req}')
        else:
            if direct or unresolved:raise ValueError(f'Aresta sem REQUER: {i}')
            requirement=None
        out.append({'id':original['id'],'paragrafo_docx':i,'grupo':'ARTISTA MARCIAL',
                    'categoria':'classe','titulo':original['titulo'],
                    'requisito_literal':req,'requisito_estruturado':requirement,
                    'partes':parts,'opcoes':opts,
                    'texto_original':'\n'.join(p['texto'] for p in parts),
                    'texto_integral_conferido':True,
                    'estado':'transcricao_integral_sem_revisao_mecanica',
                    'aplicacao_canonica':False,'origem_inventario_id':original['id']})
    counts={'total_movimentos':len(out),
            'total_com_requer':sum(bool(m['requisito_literal']) for m in out),
            'total_vinculos_exatos':sum((m['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for m in out),
            'total_requer_composto_ou_condicional':sum((m['requisito_estruturado'] or {}).get('tipo')=='todos' for m in out),
            'total_opcoes':sum(len(m['opcoes']) for m in out),
            'total_partes':sum(len(m['partes']) for m in out)}
    if tuple(counts[k] for k in ('total_movimentos','total_com_requer','total_vinculos_exatos','total_requer_composto_ou_condicional'))!=(42,19,18,1):
        raise ValueError(f'Contagens divergentes: {counts}')
    if len(occupied)!=counts['total_partes']:raise ValueError('Partes duplicadas')
    return {'schema_version':'0.1.0','fonte':inventory['fonte'],
            'fonte_sha256':source_hash,'grupo':'ARTISTA MARCIAL',
            'posicao':'indices_paragrafo_docx_zero_based',
            'descricao':'Transcrição integral dos 42 movimentos; VOO requer Forma do Ar E Corpo Sem Peso. Não homologa efeitos nem habilita aquisição.',
            'movimentos':out,'integridade':counts}

if __name__=='__main__':
    docx=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inv=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    target=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_artista_marcial/hack-carcaca-dw-artista-marcial-integrais.json')
    target.parent.mkdir(parents=True,exist_ok=True)
    result=extract(docx,inv)
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result['integridade'],ensure_ascii=False,indent=2))