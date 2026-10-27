#!/usr/bin/env python3
"""Transcrição fiel dos movimentos BARDO do Hack Carcaça DW.
Uso: python extrair_bardo_docx.py MANUSCRITO.docx INVENTARIO.json SAIDA.json
Não homologa REQUER: exige Arte Arcana e dois outros movimentos distintos de Bardo.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import hashlib, json, sys, re, unicodedata

FIRST, LAST, NEXT_HEADING = 1753, 1851, 1860
SPECIAL_COUNT = {1843: "Arte Arcana e dois outros Movimentos de Bardo"}

def slug(text):
    return re.sub(r'[^a-z0-9]+', '-', unicodedata.normalize('NFKD', text)
                  .encode('ascii', 'ignore').decode().lower()).strip('-')

def extract(docx_path, inventory_path):
    docx_path, inventory_path = Path(docx_path), Path(inventory_path)
    inventory = json.loads(inventory_path.read_text(encoding='utf8'))
    source_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if source_hash != inventory['fonte_sha256']:
        raise ValueError('SHA256 da fonte diverge do inventário de origem')
    doc = Document(docx_path)
    rows = [m for m in inventory['movimentos'] if m['grupo'] == 'BARDO']
    if len(rows) != 32 or rows[0]['paragrafo_docx'] != FIRST or rows[-1]['paragrafo_docx'] != LAST:
        raise ValueError('Quantidade ou âncoras documentais BARDO inesperadas')
    if (doc.paragraphs[1751].text.strip() != 'BARDO' or
            doc.paragraphs[NEXT_HEADING].text.strip() != 'ARTISTA MARCIAL' or
            not doc.paragraphs[NEXT_HEADING].style.name.startswith('Heading')):
        raise ValueError('Limites do grupo BARDO não confirmados')
    titles={}
    for m in rows:titles.setdefault(slug(m['titulo']), []).append(m)
    all_by_id = {m['id']: m for m in inventory['movimentos']}
    edges={m['origem_movimento_id']:m for m in inventory['dependencias_exatas_candidatas']}
    pending={m['origem_movimento_id']:m for m in inventory['dependencias_nao_resolvidas']}
    out=[];occupied=set()
    for k, original in enumerate(rows):
        i=original['paragrafo_docx']
        stop=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else NEXT_HEADING
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=original['titulo']:
            raise ValueError(f'Cabeçalho alterado em {i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=original['requisito_literal']:
            raise ValueError(f'Requisito diverge do inventário #{i}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError(f'Corpo vazio em {i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];opts=[]
        if i in occupied:raise ValueError(f'Parágrafo duplicado {i}')
        occupied.add(i)
        for j in range(i+1,stop):
            p=doc.paragraphs[j]
            if p.style.name.startswith('Heading'):raise ValueError(f'Cabeçalho inesperado em {j}')
            t=p.text.strip()
            if not t:continue
            if j in occupied:raise ValueError(f'Parte duplicada {j}')
            occupied.add(j)
            typ='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':t})
            if typ=='opcao':opts.append({'paragrafo_docx':j,'texto':t})
        direct=edges.get(original['id']);undetermined=pending.get(original['id'])
        if req:
            if bool(direct)==bool(undetermined):raise ValueError(f'REQUER mal inventariado: {i}')
            if direct:
                target=all_by_id[direct['requisito_movimento_id']]
                if slug(req)!=slug(target['titulo']):raise ValueError(f'Alvo não corresponde à literalidade: {i}')
                requirement={'tipo':'movimento_exato','movimento_id':target['id'],
                             'titulo_requerido':target['titulo'],'escopo':direct['escopo'],
                             'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            elif i in SPECIAL_COUNT:
                if req!=SPECIAL_COUNT[i] or undetermined['requisito_literal']!=req:
                    raise ValueError(f'REQUER com cardinalidade mudou no DOCX: {i}')
                matches=titles.get(slug('Arte Arcana'), [])
                if len(matches)!=1:raise ValueError('Arte Arcana ambígua ou inexistente')
                arte=matches[0]
                requirement={'tipo':'todos',
                    'itens':[
                        {'tipo':'movimento_exato','movimento_id':arte['id'],
                         'titulo_requerido':arte['titulo']},
                        {'tipo':'quantidade_movimentos_grupo',
                         'grupo':'BARDO','quantidade_minima':2,
                         'excluir_movimentos_ids':[arte['id'],original['id']],
                         'texto_literal':'dois outros Movimentos de Bardo',
                         'status':'quantidade_por_titulos_candidatos_pendente_homologacao'}],
                    'operador_literal':'e','status':'interpretacao_literal_pendente_homologacao'}
            else:
                raise ValueError(f'Referência especial inesperada: {i}: {req}')
        else:
            if original['id'] in edges or original['id'] in pending:
                raise ValueError(f'Aresta sem REQUER no cabeçalho: {i}')
            requirement=None
        out.append({'id':original['id'],'paragrafo_docx':i,'grupo':'BARDO',
                    'categoria':'classe','titulo':original['titulo'],
                    'requisito_literal':req,'requisito_estruturado':requirement,
                    'partes':parts,'opcoes':opts,
                    'texto_original':'\n'.join(x['texto'] for x in parts),
                    'texto_integral_conferido':True,
                    'estado':'transcricao_integral_sem_revisao_mecanica',
                    'aplicacao_canonica':False,'origem_inventario_id':original['id']})
    counts={'total_movimentos':len(out),
            'total_com_requer':sum(bool(m['requisito_literal']) for m in out),
            'total_vinculos_exatos':sum((m['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for m in out),
            'total_requer_composto_ou_condicional':sum((m['requisito_estruturado'] or {}).get('tipo')=='todos' for m in out),
            'total_opcoes':sum(len(m['opcoes']) for m in out),
            'total_partes':sum(len(m['partes']) for m in out)}
    if (counts['total_movimentos'],counts['total_com_requer'],counts['total_vinculos_exatos'],
        counts['total_requer_composto_ou_condicional']) != (32,5,4,1):
        raise ValueError(f'Contagens do BARDO alteradas: {counts}')
    if len(occupied)!=counts['total_partes']:raise ValueError('Índices de parágrafo em conflito')
    return {'schema_version':'0.1.0','fonte':inventory['fonte'],'fonte_sha256':source_hash,
            'grupo':'BARDO','posicao':'indices_paragrafo_docx_zero_based',
            'descricao':'Transcrição integral do Bardo. REQUER obrigatório. A Canção que Fez o Mundo exige Arte Arcana e dois movimentos distintos adicionais de Bardo. Nenhuma homologação mecânica ou aquisição automática.',
            'movimentos':out,'integridade':counts}

if __name__=='__main__':
    docx=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inventory=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    target=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_bardo/hack-carcaca-dw-bardo-integrais.json')
    target.parent.mkdir(parents=True,exist_ok=True)
    result=extract(docx,inventory)
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(result['integridade'],ensure_ascii=False,indent=2))