#!/usr/bin/env python3
"""Reextrai movimentos PALADINO do manuscrito autoral para JSON gzip.
Uso: python scripts/extrair_paladino_docx.py MANUSCRITO.docx INVENTARIO.json [SAIDA.gz]
Requer python-docx; o manuscrito e mantido fora do repositorio.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import gzip,hashlib,json,re,sys,unicodedata

def slug(s):
    return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',s)
                  .encode('ascii','ignore').decode().lower()).strip('-')

def extract(src,inventory):
    src=Path(src);inv=json.loads(Path(inventory).read_text(encoding='utf-8'))
    sha=hashlib.sha256(src.read_bytes()).hexdigest()
    if sha!=inv['fonte_sha256']:raise ValueError('Manuscrito SHA256 divergente')
    doc=Document(src)
    rows=[m for m in inv['movimentos'] if m['grupo']=='PALADINO']
    if len(rows)!=28 or rows[0]['paragrafo_docx']!=1968 or rows[-1]['paragrafo_docx']!=2048:
        raise ValueError('Inventario de PALADINO divergente')
    if doc.paragraphs[1966].text.strip()!='PALADINO' or doc.paragraphs[2054].text.strip()!='CLÉRIGO':
        raise ValueError('Limites da classe divergentes no DOCX')
    deps={d['origem_movimento_id']:d for d in inv['dependencias_exatas_candidatas']}
    pending={d['origem_movimento_id']:d for d in inv['dependencias_nao_resolvidas']}
    by_id={m['id']:m for m in inv['movimentos']}
    by_title={slug(m['titulo']):m for m in rows}
    out=[];occupied=set()
    for k,m in enumerate(rows):
        i=m['paragrafo_docx'];end=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else 2054
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=m['titulo']:raise ValueError(f'Titulo divergente #{i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=m['requisito_literal']:raise ValueError(f'REQUER divergente #{i}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError(f'Efeito vazio #{i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];options=[]
        for j in range(i+1,end):
            p=doc.paragraphs[j];txt=p.text.strip()
            if p.style.name.startswith('Heading'):raise ValueError('Heading dentro do movimento')
            if not txt:continue
            typ='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':txt})
            if typ=='opcao':options.append({'paragrafo_docx':j,'texto':txt})
        for p in parts:
            if p['paragrafo_docx'] in occupied:raise ValueError('Parte duplicada')
            occupied.add(p['paragrafo_docx'])
        if req:
            edge=deps.get(m['id'])
            if not edge or m['id'] in pending:raise ValueError('REQUER sem alvo exato')
            target=by_id[edge['requisito_movimento_id']]
            if slug(req)!=slug(target['titulo']):raise ValueError('Alvo REQUER divergente')
            structured={'tipo':'movimento_exato','movimento_id':target['id'],
              'titulo_requerido':target['titulo'],'escopo':edge['escopo'],
              'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
        else:
            if m['id'] in deps or m['id'] in pending:raise ValueError('Aresta sem REQUER')
            structured=None
        grants=[]
        if m['titulo']=='MISSÃO SAGRADA':
            full='\n'.join(p['texto'] for p in parts)
            if 'Você também recebe livremente os movimentos Juramento Quebrado e Penitência.' not in full:
                raise ValueError('Concessao gratuita nao demonstrada no DOCX')
            for name in ('JURAMENTO QUEBRADO','PENITÊNCIA'):
                target=by_title[slug(name)]
                if slug(target['requisito_literal'])!=slug(m['titulo']):
                    raise ValueError('Concessao nao dispensa REQUER')
                grants.append({'movimento_id':target['id'],'titulo':target['titulo'],
                  'concessao_literal':'recebe livremente',
                  'requer_permanece_obrigatorio':True,'aplicacao_automatica':False})
        out.append({'id':m['id'],'paragrafo_docx':i,'grupo':'PALADINO','categoria':'classe',
          'titulo':m['titulo'],'requisito_literal':req,'requisito_estruturado':structured,
          'concessoes_gratuitas_literais':grants,'partes':parts,'opcoes':options,
          'texto_original':'\n'.join(p['texto'] for p in parts),
          'texto_integral_conferido':True,'estado':'transcricao_integral_sem_revisao_mecanica',
          'aplicacao_canonica':False,'origem_inventario_id':m['id']})
    integrity={'total_movimentos':len(out),
      'total_com_requer':sum(bool(m['requisito_literal']) for m in out),
      'total_vinculos_exatos':sum((m['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for m in out),
      'total_opcoes':sum(len(m['opcoes']) for m in out),
      'total_partes':sum(len(m['partes']) for m in out),
      'total_concessoes_gratuitas_literais':sum(len(m['concessoes_gratuitas_literais']) for m in out)}
    expected={'total_movimentos':28,'total_com_requer':13,'total_vinculos_exatos':13,
      'total_opcoes':47,'total_partes':83,'total_concessoes_gratuitas_literais':2}
    if integrity!=expected or len(occupied)!=83:raise ValueError(f'Integridade divergente: {integrity}')
    return {'schema_version':'0.1.0','fonte':inv['fonte'],'fonte_sha256':sha,
      'grupo':'PALADINO','posicao':'indices_paragrafo_docx_zero_based',
      'descricao':'Transcrição integral de Paladino, REQUER literais e concessões gratuitas com dependências obrigatórias; não homologa mecânicas nem altera cânone.',
      'movimentos':out,'integridade':integrity}

if __name__=='__main__':
    if len(sys.argv)<3:raise SystemExit(__doc__)
    dest=Path(sys.argv[3] if len(sys.argv)>3 else 'data/lore/hack-carcaca-dw-paladino-integrais.json.gz')
    data=extract(sys.argv[1],sys.argv[2])
    dest.parent.mkdir(parents=True,exist_ok=True)
    dest.write_bytes(gzip.compress(json.dumps(data,ensure_ascii=False,separators=(',',':')).encode('utf-8'),mtime=0))
    print(json.dumps(data['integridade'],ensure_ascii=False,indent=2))
