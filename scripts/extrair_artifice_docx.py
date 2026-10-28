#!/usr/bin/env python3
"""Extrai ARTÍFICE do DOCX original e reproduz o arquivo de fonte gzip.
Uso: python scripts/extrair_artifice_docx.py MANUSCRITO.docx INVENTARIO.json [SAIDA.gz]
Requer python-docx e acesso ao manuscrito autoral; nao altera o canone.
"""
from __future__ import annotations
import gzip,hashlib,json,re,sys,unicodedata
from pathlib import Path
from docx import Document
def slug(s):
    return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()).strip('-')
def extract(source,inventory):
    source=Path(source);inv=json.loads(Path(inventory).read_text(encoding='utf-8'))
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    if digest!=inv['fonte_sha256']:raise ValueError('SHA-256 do manuscrito divergente')
    doc=Document(source)
    rows=[m for m in inv['movimentos'] if m['grupo']=='ARTÍFICE']
    if len(rows)!=23 or rows[0]['paragrafo_docx']!=2672 or rows[-1]['paragrafo_docx']!=2731:
        raise ValueError('Inventário de ARTÍFICE alterado')
    if doc.paragraphs[2670].text.strip()!='ARTÍFICE' or doc.paragraphs[2735].text.strip()!='MERCENÁRIO':
        raise ValueError('Limites da classe alterados')
    edges={m['origem_movimento_id']:m for m in inv['dependencias_exatas_candidatas']}
    pending={m['origem_movimento_id']:m for m in inv['dependencias_nao_resolvidas']}
    ids={m['id']:m for m in rows};moves=[];seen=set()
    for pos,m in enumerate(rows):
        i=m['paragrafo_docx'];end=rows[pos+1]['paragrafo_docx'] if pos+1<len(rows) else 2735
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=m['titulo']:raise ValueError(f'Titulo divergente {i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):req=lines.pop(0).split(':',1)[1].strip()
        if req!=m['requisito_literal']:raise ValueError(f'REQUER divergente {i}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError('Corpo vazio')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];options=[]
        for j in range(i+1,end):
            p=doc.paragraphs[j];v=p.text.strip()
            if p.style.name.startswith('Heading') and v:raise ValueError('Heading inesperado')
            if not v:continue
            kind='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':kind,'texto':v})
            if kind=='opcao':options.append({'paragrafo_docx':j,'texto':v})
        for part in parts:
            if part['paragrafo_docx'] in seen:raise ValueError('Paragrafo duplicado')
            seen.add(part['paragrafo_docx'])
        if req:
            edge=edges.get(m['id'])
            if not edge or m['id'] in pending:raise ValueError('REQUER sem alvo')
            target=ids.get(edge['requisito_movimento_id'])
            if not target or slug(req)!=slug(target['titulo']):raise ValueError('REQUER sem match literal')
            structured={'tipo':'movimento_exato','movimento_id':target['id'],
                        'titulo_requerido':target['titulo'],'escopo':edge['escopo'],
                        'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
        else:
            if m['id'] in edges or m['id'] in pending:raise ValueError('Aresta sem REQUER')
            structured=None
        moves.append({'id':m['id'],'paragrafo_docx':i,'grupo':'ARTÍFICE','categoria':'classe',
          'titulo':m['titulo'],'requisito_literal':req,'requisito_estruturado':structured,
          'partes':parts,'opcoes':options,'texto_original':'\n'.join(x['texto'] for x in parts),
          'texto_integral_conferido':True,'estado':'transcricao_integral_sem_revisao_mecanica',
          'aplicacao_canonica':False,'origem_inventario_id':m['id']})
    counts={'total_movimentos':len(moves),
      'total_com_requer':sum(bool(x['requisito_literal']) for x in moves),
      'total_vinculos_exatos':sum(bool(x['requisito_estruturado']) for x in moves),
      'total_opcoes':sum(len(x['opcoes']) for x in moves),
      'total_partes':sum(len(x['partes']) for x in moves)}
    if counts!={'total_movimentos':23,'total_com_requer':7,'total_vinculos_exatos':7,
                'total_opcoes':32,'total_partes':59} or len(seen)!=59:
        raise ValueError(f'Contagens inesperadas {counts}')
    return {'schema_version':'0.1.0','fonte':inv['fonte'],'fonte_sha256':digest,
      'grupo':'ARTÍFICE','posicao':'indices_paragrafo_docx_zero_based',
      'descricao':'Transcrição integral dos movimentos de ARTÍFICE com engenharia, alquimia, cuidados, gastronomia e requisitos literais; não homologa mecânicas.',
      'movimentos':moves,'integridade':counts}
if __name__=='__main__':
    if len(sys.argv)<3:raise SystemExit(__doc__)
    out=Path(sys.argv[3] if len(sys.argv)>3 else 'data/lore/hack-carcaca-dw-artifice-integrais.json.gz')
    data=extract(sys.argv[1],sys.argv[2])
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_bytes(gzip.compress(json.dumps(data,ensure_ascii=False,separators=(',',':')).encode(),mtime=0))
    print(data['integridade'])
