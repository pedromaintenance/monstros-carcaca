#!/usr/bin/env python3
"""Reextrai Clérigo do manuscrito DOCX, cruzando-o ao inventário original.
Uso: python scripts/extrair_clerigo_docx.py MANUSCRITO.docx INVENTARIO.json [SAIDA.gz]
Requer python-docx e manuscrito autoral, que não é distribuído neste repositório.
"""
from __future__ import annotations
from docx import Document
from pathlib import Path
import gzip,hashlib,json,re,sys,unicodedata

FIRST,LAST,NEXT_HEADING=2056,2126,2128
def slug(s):
    return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',s).encode('ASCII','ignore').decode().lower()).strip('-')

def extract(docx_path,inventory_path):
    docpath,invpath=Path(docx_path),Path(inventory_path)
    inv=json.loads(invpath.read_text(encoding='utf-8'))
    sha=hashlib.sha256(docpath.read_bytes()).hexdigest()
    if sha!=inv['fonte_sha256']:raise ValueError('Manuscrito diferente da fonte indexada')
    doc=Document(docpath)
    rows=[m for m in inv['movimentos'] if m['grupo']=='CLÉRIGO']
    if len(rows)!=37 or rows[0]['paragrafo_docx']!=FIRST or rows[-1]['paragrafo_docx']!=LAST:
        raise ValueError('Inventário Clérigo alterado')
    if doc.paragraphs[2054].text.strip()!='CLÉRIGO' or doc.paragraphs[NEXT_HEADING].text.strip()!='DRUIDA':
        raise ValueError('Fronteiras das classes alteradas no DOCX')
    deps={d['origem_movimento_id']:d for d in inv['dependencias_exatas_candidatas']}
    pending={d['origem_movimento_id']:d for d in inv['dependencias_nao_resolvidas']}
    by_id={m['id']:m for m in inv['movimentos']}
    moves=[];used=set()
    for k,m in enumerate(rows):
        i=m['paragrafo_docx'];end=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else NEXT_HEADING
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=m['titulo']:raise ValueError('Título divergente')
        req=lines.pop(0).split(':',1)[1].strip() if lines and lines[0].strip().startswith('REQUER:') else None
        if req!=m['requisito_literal']:raise ValueError('REQUER divergente')
        body='\n'.join(lines).strip()
        if not body:raise ValueError('Corpo vazio')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];options=[]
        for j in range(i+1,end):
            p=doc.paragraphs[j];txt=p.text.strip()
            if p.style.name.startswith('Heading'):raise ValueError('Heading inesperado')
            if not txt:continue
            typ='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':txt})
            if typ=='opcao':options.append({'paragrafo_docx':j,'texto':txt})
        for p in parts:
            if p['paragrafo_docx'] in used:raise ValueError('Parágrafo duplicado')
            used.add(p['paragrafo_docx'])
        if req:
            edge=deps.get(m['id'])
            if not edge or m['id'] in pending:raise ValueError('REQUER não resolvido')
            target=by_id[edge['requisito_movimento_id']]
            if slug(req)!=slug(target['titulo']):raise ValueError('Título de REQUER diferente')
            structured={'tipo':'movimento_exato','movimento_id':target['id'],
              'titulo_requerido':target['titulo'],'escopo':edge['escopo'],
              'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
        else:
            if m['id'] in deps or m['id'] in pending:raise ValueError('Dependência sem REQUER')
            structured=None
        moves.append({'id':m['id'],'paragrafo_docx':i,'grupo':'CLÉRIGO',
          'categoria':'classe','titulo':m['titulo'],'requisito_literal':req,
          'requisito_estruturado':structured,'concessoes_gratuitas_literais':[],
          'partes':parts,'opcoes':options,'texto_original':'\n'.join(p['texto'] for p in parts),
          'texto_integral_conferido':True,'estado':'transcricao_integral_sem_revisao_mecanica',
          'aplicacao_canonica':False,'origem_inventario_id':m['id']})
    integrity={'total_movimentos':len(moves),
      'total_com_requer':sum(bool(m['requisito_literal']) for m in moves),
      'total_vinculos_exatos':sum(bool(m['requisito_estruturado']) for m in moves),
      'total_opcoes':sum(len(m['opcoes']) for m in moves),
      'total_partes':sum(len(m['partes']) for m in moves),
      'total_concessoes_gratuitas_literais':0}
    if integrity!={'total_movimentos':37,'total_com_requer':19,'total_vinculos_exatos':19,
                  'total_opcoes':26,'total_partes':68,'total_concessoes_gratuitas_literais':0}:
        raise ValueError('Integridade documental diferente: '+str(integrity))
    return {'schema_version':'0.1.0','fonte':inv['fonte'],'fonte_sha256':sha,
      'grupo':'CLÉRIGO','posicao':'indices_paragrafo_docx_zero_based',
      'descricao':'Transcrição integral de Clérigo com textos, opções e REQUER diretos rastreáveis; não homologa mecânicas nem altera cânone.',
      'movimentos':moves,'integridade':integrity}

if __name__=='__main__':
    if len(sys.argv)<3:raise SystemExit(__doc__)
    dest=Path(sys.argv[3] if len(sys.argv)>3 else 'data/lore/hack-carcaca-dw-clerigo-integrais.json.gz')
    obj=extract(sys.argv[1],sys.argv[2])
    dest.parent.mkdir(parents=True,exist_ok=True)
    dest.write_bytes(gzip.compress(json.dumps(obj,ensure_ascii=False,separators=(',',':')).encode('utf-8'),mtime=0))
    print(json.dumps(obj['integridade'],ensure_ascii=False,indent=2))
