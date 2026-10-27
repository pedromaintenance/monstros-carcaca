#!/usr/bin/env python3
"""Transcrição fiel dos movimentos LADINO do Hack Carcaça DW.
Uso: python extrair_ladino_docx.py MANUSCRITO.docx INVENTARIO.json SAIDA.json
Não homologa REQUER, nem corrige referências indefinidas no texto do autor.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import hashlib, json, sys, re, unicodedata

FIRST, LAST, NEXT_HEADING = 1588, 1742, 1751
SPECIAL_OR = {1619: ("Heist ou Plano B", ["HEIST", "PLANO B"])}
SPECIAL_UNRESOLVED = {1724: "Perfil"}

def slug(text):
    return re.sub(r'[^a-z0-9]+', '-', unicodedata.normalize('NFKD', text)
                  .encode('ascii', 'ignore').decode().lower()).strip('-')

def extract(docx_path, inventory_path):
    docx_path, inventory_path = Path(docx_path), Path(inventory_path)
    inventory = json.loads(inventory_path.read_text(encoding='utf8'))
    source_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if source_hash != inventory['fonte_sha256']:
        raise ValueError('SHA256 da fonte diverge do inventário de origem')
    doc = Document(docx_path)
    rows = [m for m in inventory['movimentos'] if m['grupo'] == 'LADINO']
    if len(rows) != 52 or rows[0]['paragrafo_docx'] != FIRST or rows[-1]['paragrafo_docx'] != LAST:
        raise ValueError('Quantidade ou âncoras documentais LADINO inesperadas')
    if (doc.paragraphs[1586].text.strip() != 'LADINO' or
            doc.paragraphs[NEXT_HEADING].text.strip() != 'BARDO' or
            not doc.paragraphs[NEXT_HEADING].style.name.startswith('Heading')):
        raise ValueError('Limites do grupo LADINO não confirmados')
    titles={}
    for m in rows:titles.setdefault(slug(m['titulo']), []).append(m)
    all_by_id = {m['id']: m for m in inventory['movimentos']}
    edges={m['origem_movimento_id']:m for m in inventory['dependencias_exatas_candidatas']}
    pending={m['origem_movimento_id']:m for m in inventory['dependencias_nao_resolvidas']}
    out=[];occupied=set()
    for k, original in enumerate(rows):
        i=original['paragrafo_docx']
        stop=rows[k+1]['paragrafo_docx'] if k+1<len(rows) else NEXT_HEADING
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=original['titulo']:
            raise ValueError(f'Cabeçalho alterado em {i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=original['requisito_literal']:
            raise ValueError(f'Requisito diverge do inventário #{i}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError(f'Corpo vazio em {i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];opts=[]
        if i in occupied:raise ValueError(f'Parágrafo duplicado {i}')
        occupied.add(i)
        for j in range(i+1,stop):
            p=doc.paragraphs[j]
            if p.style.name.startswith('Heading'):raise ValueError(f'Cabeçalho inesperado em {j}')
            t=p.text.strip()
            if not t:continue
            if j in occupied:raise ValueError(f'Parte duplicada {j}')
            occupied.add(j)
            typ='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':typ,'texto':t})
            if typ=='opcao':opts.append({'paragrafo_docx':j,'texto':t})
        direct=edges.get(original['id']);undetermined=pending.get(original['id'])
        if req:
            if bool(direct)==bool(undetermined):raise ValueError(f'REQUER mal inventariado: {i}')
            if direct:
                target=all_by_id[direct['requisito_movimento_id']]
                if slug(req)!=slug(target['titulo']):raise ValueError(f'Alvo não corresponde à literalidade: {i}')
                requirement={'tipo':'movimento_exato','movimento_id':target['id'],
                             'titulo_requerido':target['titulo'],'escopo':direct['escopo'],
                             'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            elif i in SPECIAL_OR:
                expected, names=SPECIAL_OR[i]
                if req!=expected or undetermined['requisito_literal']!=req:
                    raise ValueError(f'REQUER OU mudou no DOCX: {i}')
                nodes=[]
                for name in names:
                    matches=titles.get(slug(name),[])
                    if len(matches)!=1:raise ValueError(f'Alvo ambíguo de OU: {name}')
                    nodes.append({'tipo':'movimento_exato','movimento_id':matches[0]['id'],
                                  'titulo_requerido':matches[0]['titulo']})
                requirement={'tipo':'qualquer','itens':nodes,'operador_literal':'ou',
                             'status':'interpretacao_literal_pendente_homologacao'}
            elif i in SPECIAL_UNRESOLVED:
                if req!=SPECIAL_UNRESOLVED[i] or undetermined['requisito_literal']!=req:
                    raise ValueError(f'REQUER sem alvo mudou no DOCX: {i}')
                if titles.get(slug(req),[]) or any(slug(m['titulo'])==slug(req) for m in all_by_id.values()):
                    raise ValueError('Referência Perfil agora tem candidato; realizar revisão antes de publicar')
                requirement={'tipo':'referencia_sem_alvo','texto_literal':req,
                             'candidatos_confirmados':[],
                             'status':'exige_revisao_autoral_antes_de_qualquer_aquisicao'}
            else:
                raise ValueError(f'Referência especial inesperada: {i}: {req}')
        else:
            if original['id'] in edges or original['id'] in pending:
                raise ValueError(f'Aresta sem REQUER no cabeçalho: {i}')
            requirement=None
        out.append({'id':original['id'],'paragrafo_docx':i,'grupo':'LADINO',
                    'categoria':'classe','titulo':original['titulo'],
                    'requisito_literal':req,'requisito_estruturado':requirement,
                    'partes':parts,'opcoes':opts,
                    'texto_original':'\n'.join(x['texto'] for x in parts),
                    'texto_integral_conferido':True,
                    'estado':'transcricao_integral_sem_revisao_mecanica',
                    'aplicacao_canonica':False,'origem_inventario_id':original['id']})
    counts={'total_movimentos':len(out),
            'total_com_requer':sum(bool(m['requisito_literal']) for m in out),
            'total_vinculos_exatos':sum((m['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for m in out),
            'total_requer_composto_ou_condicional':sum((m['requisito_estruturado'] or {}).get('tipo')=='qualquer' for m in out),
            'total_requer_sem_alvo':sum((m['requisito_estruturado'] or {}).get('tipo')=='referencia_sem_alvo' for m in out),
            'total_opcoes':sum(len(m['opcoes']) for m in out),
            'total_partes':sum(len(m['partes']) for m in out)}
    if (counts['total_movimentos'],counts['total_com_requer'],counts['total_vinculos_exatos'],
        counts['total_requer_composto_ou_condicional'],counts['total_requer_sem_alvo']) != (52,6,4,1,1):
        raise ValueError(f'Contagens do LADINO alteradas: {counts}')
    if len(occupied)!=counts['total_partes']:raise ValueError('Índices de parágrafo em conflito')
    return {'schema_version':'0.1.0','fonte':inventory['fonte'],'fonte_sha256':source_hash,
            'grupo':'LADINO','posicao':'indices_paragrafo_docx_zero_based',
            'descricao':'Transcrição integral do Ladino. REQUER é obrigatório; "Perfil" permanece sem alvo documentado e bloqueia liberação automática. Não há homologação mecânica.',
            'movimentos':out,'integridade':counts}

if __name__=='__main__':
    docx=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inventory=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    target=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_ladino/hack-carcaca-dw-ladino-integrais.json')
    target.parent.mkdir(parents=True,exist_ok=True)
    result=extract(docx,inventory)
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(result['integridade'],ensure_ascii=False,indent=2))