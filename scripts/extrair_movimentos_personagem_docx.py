#!/usr/bin/env python3
"""Extrai movimentos de povos e catálogo preliminar de títulos de classe do DOCX autoral.
Uso: python extrair_proximo_ciclo.py [CAMINHO_DOCX] [DIR_SAIDA]
Mantém índices zero-based de parágrafos e não infere REQUER nem produz fichas jogáveis.
"""
from __future__ import annotations
from collections import Counter, defaultdict
from docx import Document
import hashlib, json, re, sys, unicodedata
from pathlib import Path

RACE_START, CLASS_START = 824, 1044
RACE_TAGS = {'Cultura','Corpo','Natureza','Corpo/Cultura','Tradição','Natureza/Cultura',
             'Cultura/Magia','Memória Cultural','Fé','Fé/Cultura','Magia/Cultura'}

def slug(t):
    t=unicodedata.normalize('NFKD', t).encode('ASCII','ignore').decode('ASCII').lower()
    return re.sub(r'[^a-z0-9]+','-',t).strip('-')

def heading(t):
    f=t.strip().split('\n')[0].strip()
    if len(f)>90 or not f:
        return False
    n=sum(c.isalpha() for c in f)
    return n>=2 and sum(c.isupper() for c in f)==n

def extract(source):
    doc=Document(source)
    assert len(doc.tables)==33
    sha=hashlib.sha256(Path(source).read_bytes()).hexdigest()
    races, classes=[],[]
    group=None; current=None
    for i in range(RACE_START+1,CLASS_START):
        p=doc.paragraphs[i];t=p.text.strip()
        if p.style.name=='Heading 3':
            group=t;current=None;continue
        if not t:continue
        first=t.split('\n')[0].strip()
        title_marker=first.rsplit(' — ',1) if ' — ' in first else None
        if p.style.name!='List Paragraph' and title_marker and title_marker[1] in RACE_TAGS:
            title, marker=title_marker
            current={'id':f'movimento-povo:paragrafo-{i}','paragrafo_docx':i,
                     'grupo':group,'categoria':'povo','titulo':title,'marcador_original':marker,
                     'texto_primeiro_paragrafo':t.split('\n',1)[1].strip() if '\n' in t else '',
                     'partes':[], 'opcoes':[], 'requisito_literal':None,
                     'requisito_auditado':'sem_requer_no_cabecalho',
                     'estado':'transcricao_integral_sem_revisao_mecanica',
                     'aplicacao_canonica':False}
            if not current['texto_primeiro_paragrafo']:
                raise ValueError(f'Sem corpo no título {i}: {title}')
            current['partes'].append({'paragrafo_docx':i,'tipo':'corpo','texto':current['texto_primeiro_paragrafo']})
            races.append(current);continue
        if current is None:raise ValueError(f'Parágrafo raça fora de movimento: {i} {t[:80]}')
        kind='opcao' if p.style.name=='List Paragraph' else 'continuacao'
        current['partes'].append({'paragrafo_docx':i,'tipo':kind,'texto':t})
        if kind=='opcao':current['opcoes'].append({'paragrafo_docx':i,'texto':t})
    for x in races:
        x['texto_original']='\n'.join(part['texto'] for part in x['partes'])
        del x['texto_primeiro_paragrafo']
    group=None; current=None
    for i in range(CLASS_START+1,len(doc.paragraphs)):
        p=doc.paragraphs[i];t=p.text.strip()
        if p.style.name=='Heading 3':group=t;current=None;continue
        if not t:continue
        first=t.split('\n')[0].strip()
        if p.style.name!='List Paragraph' and heading(t):
            raw=t.splitlines();requer=[x.strip().split(':',1)[1].strip() for x in raw[1:] if x.strip().startswith('REQUER:')]
            if len(requer)>1: raise ValueError(f'Mais de um REQUER no título {i}')
            current={'id':f'movimento-classe:paragrafo-{i}', 'paragrafo_docx':i,
                     'grupo':group,'categoria':'classe','titulo':first,
                     'requisito_literal':requer[0] if requer else None,
                     'estado':'titulo_requer_autoextraidos_pendentes_revisao',
                     'texto_integral_curado':False,'aplicacao_canonica':False}
            classes.append(current);continue
        if current is None: raise ValueError(f'Parágrafo de classe sem movimento inicial: {i} {t[:70]}')
    # Link only unambiguous literal requirements (no ``qualquer`` or conjunction).
    lookup=defaultdict(list)
    for x in classes: lookup[slug(x['titulo'])].append(x)
    edges=[];unresolved=[]
    for x in classes:
        literal=x['requisito_literal']
        if not literal:continue
        same=[m for m in lookup[slug(literal)] if m['grupo']==x['grupo']]
        cand=same if same else lookup[slug(literal)]
        if len(cand)==1 and cand[0]['id']!=x['id']:
            edges.append({'id':f'requer:{x["paragrafo_docx"]}-{cand[0]["paragrafo_docx"]}',
                'origem_movimento_id':x['id'],'requisito_movimento_id':cand[0]['id'],
                'requisito_literal':literal,'escopo':'mesmo_grupo' if same else 'outro_grupo',
                'status':'vinculo_exato_a_confirmar','aplicacao_canonica':False})
        else:
            unresolved.append({'origem_movimento_id':x['id'],'grupo':x['grupo'],
                'titulo':x['titulo'],'requisito_literal':literal,
                'candidatos_ids':[m['id'] for m in cand],
                'motivo':'multiplos_candidatos' if len(cand)>1 else 'nao_corresponde_a_um_titulo_isolado',
                'status':'exige_analise_editorial','aplicacao_canonica':False})
    r={'schema_version':'0.1.0','fonte':'Hack Carcaça Dungeon World(1).docx',
       'fonte_sha256':sha,'posicao':'indices_paragrafo_docx_zero_based',
       'descricao':'Movimentos de raça/povo transcritos por parágrafos DOCX incluindo opções e continuações, sem revisão mecânica.',
       'movimentos':races}
    c={'schema_version':'0.1.0','fonte':'Hack Carcaça Dungeon World(1).docx',
       'fonte_sha256':sha,'posicao':'indices_paragrafo_docx_zero_based',
       'descricao':'Inventário preliminar de títulos e pré-requisitos literais de classe. Não confundir com texto integral curado.',
       'movimentos':classes,'dependencias_exatas_candidatas':edges,'dependencias_nao_resolvidas':unresolved}
    return r,c

if __name__=='__main__':
    source=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    out=Path(sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos');out.mkdir(exist_ok=True,parents=True)
    race,cl=extract(source)
    for p,obj in [('hack-carcaca-dw-movimentos-povos-integrais.json',race),('hack-carcaca-dw-inventario-classes-requer.json',cl)]:
        (out/p).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print('SHA',race['fonte_sha256'])
    print('RACES',len(race['movimentos']),Counter(m['grupo'] for m in race['movimentos']),'OPTIONS',sum(len(m['opcoes']) for m in race['movimentos']))
    print('CLASSES',len(cl['movimentos']),Counter(m['grupo'] for m in cl['movimentos']))
    print('REQ',sum(m['requisito_literal'] is not None for m in cl['movimentos']),'EDGES',len(cl['dependencias_exatas_candidatas']),'UNRESOLVED',len(cl['dependencias_nao_resolvidas']))
    print('UNRESOLVED EXAMPLES',[(e['titulo'],e['requisito_literal'],e['motivo']) for e in cl['dependencias_nao_resolvidas'][:30]])
    print('SIZES',[(p.name,p.stat().st_size) for p in out.iterdir()])
