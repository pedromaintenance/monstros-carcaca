#!/usr/bin/env python3
"""Extrai quadros gerais diretamente das células DOCX, sem interpretar regras.
Uso: python scripts/extrair_quadros_hack.py "Hack Carcaça Dungeon World.docx" -o saida.json
Requer python-docx.
"""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document

def extrair(caminho):
    caminho = Path(caminho)
    documento = Document(caminho)
    if len(documento.tables) != 33:
        raise ValueError(f"Esperadas 33 tabelas; encontradas {len(documento.tables)}")
    quadros = []
    for indice in range(12, 33):
        tabela = documento.tables[indice]
        blocos = [
            {"linha": numero, "texto": linha.cells[0].text.strip()}
            for numero, linha in enumerate(tabela.rows)
        ]
        if not blocos[0]["texto"] or any(not bloco["texto"] for bloco in blocos):
            raise ValueError(f"Bloco vazio na tabela {indice}")
        quadros.append({
            "id": f"quadro-geral-docx:{indice}",
            "tabela_docx_indice": indice,
            "titulo": blocos[0]["texto"],
            "blocos": blocos,
            "texto_integral_extraido": True,
            "aplicacao_canonica": False,
        })
    assert len(quadros) == 21
    assert quadros[18]["blocos"][2]["texto"] == "MANDADOS PENDENTES"
    assert "LEALDADE DOS SERVOS" in "\n".join(
        bloco["texto"] for bloco in quadros[-1]["blocos"]
    )
    return {
        "schema_version": "0.1.0",
        "origem": "Hack Carcaça Dungeon World.docx",
        "origem_sha256": hashlib.sha256(caminho.read_bytes()).hexdigest(),
        "metodo": "python-docx; leitura das células em ordem documental",
        "observacao": "Extração integral bruta, não normalizada; não altera cânone.",
        "quadros": quadros,
    }

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    parser.add_argument("-o", "--output", default="quadros_gerais_auditoria5.json")
    args = parser.parse_args()
    resultado = extrair(args.docx)
    Path(args.output).write_text(
        json.dumps(resultado, ensure_ascii=False, indent=2) + "\n",
        encoding="utf8",
    )
    print(f"Extraídos {len(resultado['quadros'])} quadros")

if __name__ == "__main__":
    main()
