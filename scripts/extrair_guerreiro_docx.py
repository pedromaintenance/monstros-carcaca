#!/usr/bin/env python3
"""Extrai integralmente os movimentos de Guerreiro do Hack Carcaça DW.

Exige python-docx. Usa inventário pré-existente dos movimentos de classe para
não adivinhar títulos nem reconhecer falsos cabeçalhos nas opções.
Nenhuma regra é homologada/alterada automaticamente.
"""
from pathlib import Path
from docx import Document
import hashlib, json, sys, re, unicodedata
from collections import Counter

def slug(t):
    t=unicodedata.normalize('NFKD',t).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',t).strip('-')

def extract(docx_path,inventory_path):
    docx_path=Path(docx_path);inventory_path=Path(inventory_path)
    doc=Document(docx_path)
    if len(doc.tables)!=33 or doc.paragraphs[1045].text.strip()!="GUERREIRO" or doc.paragraphs[1291].text.strip()!="BÁRBARO":
        raise ValueError("As fronteiras de classe ou o número de tabelas do DOCX mudou")
    inventory=json.loads(inventory_path.read_text(encoding='utf8'))
    source_hash=hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if source_hash!=inventory['fonte_sha256']:
        raise ValueError('Hash da fonte não confere com o inventário anterior')
    warriors=[r for r in inventory['movimentos'] if r['grupo']=='GUERREIRO']
    if len(warriors)!=77: raise ValueError(f'Guerreiro: 77 esperados, {len(warriors)} encontrados')
    if warriors[0]['paragrafo_docx']!=1047 or warriors[-1]['paragrafo_docx']!=1289:
        raise ValueError('Âncoras alteradas')
    all_index={x['id']:x for x in inventory['movimentos']}
    edges={e['origem_movimento_id']:e for e in inventory['dependencias_exatas_candidatas']}
    unresolved={u['origem_movimento_id']:u for u in inventory['dependencias_nao_resolvidas']}
    by_title={slug(x['titulo']):x for x in warriors}
    movements=[]
    for k,original in enumerate(warriors):
        i=original['paragrafo_docx'];end=warriors[k+1]['paragrafo_docx'] if k+1<len(warriors) else 1291
        paragraph=doc.paragraphs[i];lines=paragraph.text.strip().splitlines()
        title=lines.pop(0).strip()
        if title!=original['titulo']:
            raise ValueError(f'Cabeçalho #{i} divergente: {title!r} vs {original["titulo"]!r}')
        requirement=None
        if lines and lines[0].strip().startswith('REQUER:'):
            requirement=lines.pop(0).strip().split(':',1)[1].strip()
        if requirement!=original['requisito_literal']:
            raise ValueError(f'REQUER divergente #{i}: {requirement!r} / {original["requisito_literal"]!r}')
        body='\n'.join(lines).strip()
        if not body: raise ValueError(f'Sem corpo em #{i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}]
        options=[]
        for j in range(i+1,end):
            nextp=doc.paragraphs[j]
            if nextp.style.name.startswith('Heading'):
                raise ValueError(f'Cabeçalho inesperado entre movimentos em {j}')
            text=nextp.text.strip()
            if not text:continue
            typ='opcao' if nextp.style.name=='List Paragraph' else 'continuacao'
            part={'paragrafo_docx':j,'tipo':typ,'texto':text}
            parts.append(part)
            if typ=='opcao':options.append({'paragrafo_docx':j,'texto':text})
        link=edges.get(original['id']); pend=unresolved.get(original['id'])
        if requirement:
            if bool(link)==bool(pend):raise ValueError(f'Atribuição de prereq ambígua #{i}')
            if link:
                target=all_index[link['requisito_movimento_id']]
                if slug(requirement)!=slug(target['titulo']):raise ValueError(f'Vinculo não literal #{i}')
                req={'tipo':'movimento_exato','movimento_id':target['id'],
                     'titulo_requerido':target['titulo'],'escopo':link['escopo'],
                     'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            else:
                if i==1228:
                    req={'tipo':'categoria','categoria':'movimento_de_magia',
                         'minimo':1,'status':'interpretacao_literal_pendente_homologacao'}
                elif i==1229:
                    req={'tipo':'todos','itens':[
                        {'tipo':'movimento_exato','movimento_id':by_title[slug('Surto de Ação')]['id']},
                        {'tipo':'categoria','categoria':'movimento_de_magia','minimo':1}],
                         'status':'interpretacao_literal_pendente_homologacao'}
                elif i==1231:
                    req={'tipo':'movimento_condicionado',
                         'movimento_id':by_title[slug('Arma Favorita')]['id'],
                         'condicao_literal':'aplicada a um arco ou arma de projétil',
                         'status':'interpretacao_literal_pendente_homologacao'}
                else:
                    raise ValueError(f'Requisito especial não previsto no Guerreiro: {i} {requirement}')
        else:req=None
        mov={
            'id':original['id'],'paragrafo_docx':i,'grupo':'GUERREIRO',
            'categoria':'classe','titulo':title,'requisito_literal':requirement,
            'requisito_estruturado':req,
            'partes':parts,'opcoes':options,
            'texto_original':'\n'.join(x['texto'] for x in parts),
            'texto_integral_conferido':True,
            'estado':'transcricao_integral_sem_revisao_mecanica',
            'aplicacao_canonica':False,
            'origem_inventario_id':original['id'],
        }
        movements.append(mov)
    if movements[-1]['titulo']!='CONHEÇO A ARENA': raise ValueError('Movimento final errado')
    if len(movements)!=77 or sum(bool(x['requisito_literal']) for x in movements)!=30:
        raise ValueError('Cobertura Guerreiro incorreta')
    if sum(x['requisito_estruturado']['tipo']=='movimento_exato' for x in movements if x['requisito_estruturado'])!=27:
        raise ValueError('Quantidade de vínculos exatos diferente da esperada')
    return {'schema_version':'0.1.0','fonte':'Hack Carcaça Dungeon World(1).docx',
      'fonte_sha256':source_hash,'grupo':'GUERREIRO',
      'posicao':'indices_paragrafo_docx_zero_based',
      'descricao':'Transcrição integral dos 77 movimentos de Guerreiro, incluindo opções e continuações de células de parágrafos. Dependências são literais ou candidatas não homologadas.',
      'movimentos':movements,
      'integridade':{'total_movimentos':77,'total_com_requer':30,'total_vinculos_exatos':27,
          'total_requer_composto_ou_condicional':3,
          'total_opcoes':sum(len(x['opcoes']) for x in movements),
          'total_partes':sum(len(x['partes']) for x in movements)}
    }

if __name__=='__main__':
  if len(sys.argv)<2:
      raise SystemExit('Uso: python scripts/extrair_guerreiro_docx.py <caminho_docx> [inventario_json] [saida_json]')
  root=Path(__file__).resolve().parents[1]
  d=sys.argv[1]
  inventory=sys.argv[2] if len(sys.argv)>2 else root/'data/lore/hack-carcaca-dw-inventario-classes-requer.json'
  output=Path(sys.argv[3]) if len(sys.argv)>3 else root/'data/lore/hack-carcaca-dw-guerreiro-integrais.json'
  result=extract(d,inventory)
  output.parent.mkdir(parents=True,exist_ok=True)
  output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\\n',encoding='utf8')
  print('TOTAL:',result['integridade'])
  print('SIZE:',output.stat().st_size)
  print('REQ especiais:',[(x['titulo'],x['requisito_literal']) for x in result['movimentos'] if x['requisito_estruturado'] and x['requisito_estruturado']['tipo']!='movimento_exato'])
