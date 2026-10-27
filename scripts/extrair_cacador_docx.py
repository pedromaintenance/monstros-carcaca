#!/usr/bin/env python3
"""Extração integral e determinística do CAÇADOR para o atlas Carcaça.
Uso: python extrair_cacador_docx.py <hack.docx> <inventario.json> <saida.json>
REQUER é dependência obrigatória; correspondência exata não implica homologação.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import hashlib, json, sys, re, unicodedata

FIRST, LAST, NEXT_HEADING = 1455, 1582, 1586
SPECIAL = {
    1492: ('qualquer', 'Caçador de Monstros', 'Um movimento que tenha revelado uma vulnerabilidade'),
    1520: ('todos', 'Caminhante dos Horizontes', 'Marca do Caçador'),
    1535: ('qualquer', 'Flecha de Sangue', 'Marca do Caçador'),
}

def slug(t):
    return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',t)
                  .encode('ascii','ignore').decode().lower()).strip('-')

def extract(docx_path, inventory_path):
    docx_path, inventory_path = Path(docx_path), Path(inventory_path)
    inventory = json.loads(inventory_path.read_text(encoding='utf8'))
    source_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if source_hash != inventory['fonte_sha256']:
        raise ValueError('SHA256 do manuscrito diverge do inventário de origem')
    doc = Document(docx_path)
    rows = [m for m in inventory['movimentos'] if m['grupo']=='CAÇADOR']
    if len(rows)!=53 or rows[0]['paragrafo_docx']!=FIRST or rows[-1]['paragrafo_docx']!=LAST:
        raise ValueError('Quantidade ou âncoras de Caçador inesperadas')
    if doc.paragraphs[NEXT_HEADING].text.strip() != 'LADINO' or not doc.paragraphs[NEXT_HEADING].style.name.startswith('Heading'):
        raise ValueError('Fim do grupo CAÇADOR não confirmado')
    all_index = {m['id']:m for m in inventory['movimentos']}
    edges = {m['origem_movimento_id']:m for m in inventory['dependencias_exatas_candidatas']}
    pending = {m['origem_movimento_id']:m for m in inventory['dependencias_nao_resolvidas']}
    titles = {}
    for m in rows:
        titles.setdefault(slug(m['titulo']), []).append(m)
    movements=[]; occupied=set()
    for k,original in enumerate(rows):
        i=original['paragrafo_docx']; stop = rows[k+1]['paragrafo_docx'] if k+1<len(rows) else NEXT_HEADING
        if i in occupied: raise ValueError('Índice já consumido: '+str(i))
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=original['titulo']:
            raise ValueError('Cabeçalho divergente em '+str(i))
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=original['requisito_literal']:
            raise ValueError(f'REQUER divergente #{i}: {req} / {original["requisito_literal"]}')
        body='\n'.join(lines).strip()
        if not body: raise ValueError('Sem corpo em '+str(i))
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];options=[]
        occupied.add(i)
        for j in range(i+1,stop):
            p=doc.paragraphs[j]; text=p.text.strip()
            if p.style.name.startswith('Heading'):
                raise ValueError('Outro cabeçalho dentro do movimento em '+str(j))
            if not text: continue
            if j in occupied:raise ValueError('Parágrafo duplicado: '+str(j))
            occupied.add(j)
            kind='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':kind,'texto':text})
            if kind=='opcao':options.append({'paragrafo_docx':j,'texto':text})
        direct=edges.get(original['id']);unknown=pending.get(original['id'])
        if req:
            if bool(direct)==bool(unknown):raise ValueError('REQUER não unívoco: '+original['titulo'])
            if direct:
                target=all_index[direct['requisito_movimento_id']]
                if slug(req)!=slug(target['titulo']):raise ValueError('Vínculo não literal em '+original['titulo'])
                structured={'tipo':'movimento_exato','movimento_id':target['id'],
                            'titulo_requerido':target['titulo'],'escopo':direct['escopo'],
                            'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
            else:
                if i not in SPECIAL:raise ValueError('REQUER especial não previsto: '+str(i))
                op,*required=SPECIAL[i]
                if unknown['requisito_literal']!=req:raise ValueError('REQUER especial divergente')
                if i==1492:
                    if req!='Caçador de Monstros ou Um movimento que tenha revelado uma vulnerabilidade':
                        raise ValueError('Texto composto alterado em 1492')
                elif slug(req)!=slug(' e '.join(required) if op=='todos' else ' ou '.join(required)):
                    raise ValueError('Conjunção literal divergente: '+str(i))
                nodes=[]
                for name in required:
                    targets=titles.get(slug(name),[])
                    if len(targets)==1:
                        nodes.append({'tipo':'movimento_exato','movimento_id':targets[0]['id'],
                                      'titulo_requerido':targets[0]['titulo']})
                    else:
                        # A segunda alternativa do movimento 1492 é ficcional, NÃO nome próprio.
                        if i!=1492 or name!=SPECIAL[1492][2] or targets:
                            raise ValueError('Condição sem referência controlada '+name)
                        nodes.append({'tipo':'condicao_ficcional','texto_literal':name,
                                      'estado':'exige_comprovacao_pelo_mestre'})
                structured={'tipo':op,'itens':nodes,'operador_literal':'e' if op=='todos' else 'ou',
                            'status':'interpretacao_literal_pendente_homologacao'}
        else:
            if original['id'] in edges or original['id'] in pending:
                raise ValueError('Dependência sem REQUER literal em '+str(i))
            structured=None
        movements.append({'id':original['id'],'paragrafo_docx':i,'grupo':'CAÇADOR',
                          'categoria':'classe','titulo':original['titulo'],
                          'requisito_literal':req,'requisito_estruturado':structured,
                          'partes':parts,'opcoes':options,
                          'texto_original':'\n'.join(x['texto'] for x in parts),
                          'texto_integral_conferido':True,
                          'estado':'transcricao_integral_sem_revisao_mecanica',
                          'aplicacao_canonica':False,'origem_inventario_id':original['id']})
    countreq=sum(bool(x['requisito_literal']) for x in movements)
    countdirect=sum((x['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for x in movements)
    countspecial=sum((x['requisito_estruturado'] or {}).get('tipo') in ('todos','qualquer') for x in movements)
    if (len(movements), countreq,countdirect,countspecial)!=(53,15,12,3):
        raise ValueError(f'Contagens incorretas {len(movements),countreq,countdirect,countspecial}')
    return {'schema_version':'0.1.0','fonte':inventory['fonte'],
            'fonte_sha256':source_hash,'grupo':'CAÇADOR','posicao':'indices_paragrafo_docx_zero_based',
            'descricao':'Transcrição integral literal de Caçador, opções e continuações. Dependências exatas e compostas ainda não homologadas; requisitos ficcionais não são ignorados.',
            'movimentos':movements,
            'integridade':{'total_movimentos':53,'total_com_requer':countreq,
                          'total_vinculos_exatos':countdirect,
                          'total_requer_composto_ou_condicional':countspecial,
                          'total_opcoes':sum(len(x['opcoes']) for x in movements),
                          'total_partes':sum(len(x['partes']) for x in movements)}}

if __name__=='__main__':
    docx=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inventory=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    target=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_cacador/hack-carcaca-dw-cacador-integrais.json')
    target.parent.mkdir(parents=True,exist_ok=True)
    result=extract(docx,inventory)
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(result['integridade'],ensure_ascii=False,indent=2))
