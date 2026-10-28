#!/usr/bin/env python3
"""Extrai as nove entradas finais de Escriba Yokai, Nobre e Arqueólogo do DOCX.
Uso: python scripts/extrair_classes_finais_docx.py manuscrito.docx inventario.json diretorio
Requer python-docx; o DOCX autoral não é distribuído no repositório.
"""
import sys,json,hashlib
from pathlib import Path
from docx import Document

GROUPS={
  "ESCRIBA YOKAI":(2812,2814,"escriba-yokai",2811),
  "NOBRE":(2816,2818,"nobre",2815),
  "ARQUEÓLOGO":(2820,2822,"arqueologo",2819),
}
def extract(source,inventory):
    src=Path(source);inv=json.loads(Path(inventory).read_text(encoding="utf-8"))
    digest=hashlib.sha256(src.read_bytes()).hexdigest()
    if digest!=inv["fonte_sha256"]:raise ValueError("SHA256 do manuscrito difere do inventário")
    doc=Document(src);out={}
    for group,(first,last,slug,heading) in GROUPS.items():
        if doc.paragraphs[heading].text.strip()!=group or not doc.paragraphs[heading].style.name.startswith("Heading"):
            raise ValueError("Cabeçalho de classe alterado: "+group)
        rows=[m for m in inv["movimentos"] if m["grupo"]==group]
        if [r["paragrafo_docx"] for r in rows]!=list(range(first,last+1)):
            raise ValueError("Inventário de grupo alterado: "+group)
        moves=[]
        for row in rows:
            i=row["paragrafo_docx"];p=doc.paragraphs[i]
            lines=p.text.strip().splitlines()
            if p.style.name!="Normal" or len(lines)<2 or lines[0].strip()!=row["titulo"]:
                raise ValueError("Título diferente do original: "+str(i))
            if row["requisito_literal"] is not None or any(line.strip().startswith("REQUER:") for line in lines):
                raise ValueError("REQUER inesperado no grupo "+group)
            body="\n".join(lines[1:]).strip()
            if not body:raise ValueError("Movimento sem corpo "+str(i))
            moves.append({"id":row["id"],"paragrafo_docx":i,"grupo":group,
              "categoria":"classe","titulo":row["titulo"],
              "requisito_literal":None,"requisito_estruturado":None,
              "partes":[{"paragrafo_docx":i,"tipo":"corpo","texto":body}],
              "opcoes":[],"texto_original":body,"texto_integral_conferido":True,
              "estado":"transcricao_integral_sem_revisao_mecanica",
              "aplicacao_canonica":False,"origem_inventario_id":row["id"]})
        out[slug]={"schema_version":"0.1.0","fonte":inv["fonte"],"fonte_sha256":digest,
          "grupo":group,"posicao":"indices_paragrafo_docx_zero_based",
          "descricao":f"Transcrição completa de movimentos do grupo {group}; sem homologação mecânica.",
          "movimentos":moves,"integridade":{"total_movimentos":3,"total_com_requer":0,
            "total_vinculos_exatos":0,"total_opcoes":0,"total_partes":3}}
    return out

if __name__=="__main__":
    if len(sys.argv)!=4:raise SystemExit(__doc__)
    out=Path(sys.argv[3]);out.mkdir(parents=True,exist_ok=True)
    for slug,obj in extract(sys.argv[1],sys.argv[2]).items():
        p=out/("hack-carcaca-dw-"+slug+"-integrais.json")
        p.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print(p, len(obj["movimentos"]))
