#!/usr/bin/env python3
"""Transcrição integral e verificável dos movimentos MERCENÁRIO do manuscrito.
Uso: python extrair_mercenario_docx.py MANUSCRITO.docx INVENTARIO.json [SAIDA.json]
Requer python-docx. Não homologa movimentos ou altera fichas.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import gzip, hashlib, json, re, sys, unicodedata

FIRST,LAST,HEADING,NEXT_HEADING=2736,2810,2735,2811

def slug(text):
    return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',text)
                  .encode('ASCII','ignore').decode().lower()).strip('-')

def extract(source,inventory):
    source,inventory=Path(source),Path(inventory)
    inv=json.loads(inventory.read_text(encoding='utf-8'))
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    if digest!=inv['fonte_sha256']:
        raise ValueError('SHA-256 do manuscrito divergente do inventário')
    doc=Document(source)
    group=[m for m in inv['movimentos'] if m['grupo']=='MERCENÁRIO']
    if len(group)!=31 or group[0]['paragrafo_docx']!=FIRST or group[-1]['paragrafo_docx']!=LAST:
        raise ValueError('Limites do inventário MERCENÁRIO divergentes')
    if doc.paragraphs[HEADING].text.strip()!='MERCENÁRIO' or doc.paragraphs[NEXT_HEADING].text.strip()!='ESCRIBA YOKAI':
        raise ValueError('Heading do manuscrito mudou')
    edges={x['origem_movimento_id']:x for x in inv['dependencias_exatas_candidatas']}
    pending={x['origem_movimento_id']:x for x in inv['dependencias_nao_resolvidas']}
    by_id={m['id']:m for m in group}
    moves=[];seen=set()
    for pos,original in enumerate(group):
        i=original['paragrafo_docx']
        stop=group[pos+1]['paragrafo_docx'] if pos+1<len(group) else NEXT_HEADING
        lines=doc.paragraphs[i].text.strip().splitlines()
        if not lines or lines.pop(0).strip()!=original['titulo']:
            raise ValueError(f'Título divergente no parágrafo {i}')
        req=None
        if lines and lines[0].strip().startswith('REQUER:'):
            req=lines.pop(0).split(':',1)[1].strip()
        if req!=original['requisito_literal']:
            raise ValueError(f'REQUER divergente no parágrafo {i}')
        body='\n'.join(lines).strip()
        if not body:raise ValueError(f'Movimento sem corpo: {i}')
        parts=[{'paragrafo_docx':i,'tipo':'corpo','texto':body}];options=[]
        for j in range(i+1,stop):
            p=doc.paragraphs[j];txt=p.text.strip()
            if p.style.name.startswith('Heading') and txt:
                raise ValueError(f'Cabeçalho inesperado dentro do movimento: {j}')
            if not txt:continue
            kind='opcao' if p.style.name=='List Paragraph' else 'continuacao'
            parts.append({'paragrafo_docx':j,'tipo':kind,'texto':txt})
            if kind=='opcao':options.append({'paragrafo_docx':j,'texto':txt})
        for part in parts:
            if part['paragrafo_docx'] in seen:raise ValueError(f'Parágrafo duplicado: {part}')
            seen.add(part['paragrafo_docx'])
        if not req:
            if original['id'] in edges or original['id'] in pending:
                raise ValueError(f'Aresta sem REQUER: {i}')
            structured=None
        elif original['id'] in edges:
            edge=edges[original['id']]
            target=by_id.get(edge['requisito_movimento_id'])
            if not target or slug(req)!=slug(target['titulo']):
                raise ValueError(f'REQUER sem alvo literal interno: {i}')
            structured={'tipo':'movimento_exato','movimento_id':target['id'],
                        'titulo_requerido':target['titulo'],'escopo':edge['escopo'],
                        'status':'correspondencia_exata_na_fonte_pendente_homologacao'}
        else:
            raise ValueError(f'REQUER não resolvido, não criar regra por inferência: {i} {req}')
        moves.append({'id':original['id'],'paragrafo_docx':i,'grupo':'MERCENÁRIO',
                      'categoria':'classe','titulo':original['titulo'],
                      'requisito_literal':req,'requisito_estruturado':structured,
                      'partes':parts,'opcoes':options,
                      'texto_original':'\n'.join(p['texto'] for p in parts),
                      'texto_integral_conferido':True,
                      'estado':'transcricao_integral_sem_revisao_mecanica',
                      'aplicacao_canonica':False,'origem_inventario_id':original['id']})
    stats={'total_movimentos':len(moves),
           'total_com_requer':sum(bool(m['requisito_literal']) for m in moves),
           'total_vinculos_exatos':sum((m['requisito_estruturado'] or {}).get('tipo')=='movimento_exato' for m in moves),
           'total_opcoes':sum(len(m['opcoes']) for m in moves),
           'total_partes':sum(len(m['partes']) for m in moves)}
    if len(seen)!=stats['total_partes'] or (stats['total_movimentos'],stats['total_com_requer'])!=(31,11):
        raise ValueError(f'Integridade documental incorreta: {stats}')
    return {'schema_version':'0.1.0','fonte':inv['fonte'],'fonte_sha256':digest,
            'grupo':'MERCENÁRIO','posicao':'indices_paragrafo_docx_zero_based',
            'descricao':'Transcrição integral dos movimentos de MERCENÁRIO: comandos, contratos, marinharia, pirataria, planejamento e requisitos literais; não homologa mecânicas.',
            'movimentos':moves,'integridade':stats}

if __name__=='__main__':
    manuscript=sys.argv[1] if len(sys.argv)>1 else '/mnt/data/Hack Carcaça Dungeon World(1).docx'
    inventory=sys.argv[2] if len(sys.argv)>2 else '/mnt/data/ciclo_movimentos/hack-carcaca-dw-inventario-classes-requer.json'
    dest=Path(sys.argv[3] if len(sys.argv)>3 else '/mnt/data/ciclo_mercenario/hack-carcaca-dw-mercenario-integrais.json')
    obj=extract(manuscript,inventory)
    dest.parent.mkdir(parents=True,exist_ok=True)
    dest.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    dest.with_suffix('.json.gz').write_bytes(gzip.compress(json.dumps(obj,ensure_ascii=False,separators=(',',':')).encode(),mtime=0))
    print(json.dumps(obj['integridade'],ensure_ascii=False,indent=2))
    print(f'JSON {dest.stat().st_size}, gzip {dest.with_suffix(".json.gz").stat().st_size} bytes')